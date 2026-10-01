"""
Player spreadsheets (roadmap A11): scenarios and feedback in a workbook anyone can fill in.

Players don't need Python or YAML. They fill in a workbook (Google Sheets or Excel) with tabs for the battle,
its units, the orders round by round, what they expect to happen, and, after seeing the engine's results, how
it felt. ``read_workbook`` turns a filled-in workbook into a scenario (checked, with problems pointing at the tab
and row); ``write_template`` makes a blank template, or one filled from an existing scenario for players to
adjust. The words players use (unit types, weapons, actions, ammo, feedback scales) live in config/sheets.yaml.

Not carried by the sheet: unit front sizes and free-form unit meta, and parameter overrides.
"""

import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from imperial_generals import validation
from imperial_generals.battles.setup import SetupError, battle_from_setup
from imperial_generals.config import get_config
from imperial_generals.utils.unit_types import UNIT_SUBTYPES

BATTLE_FIELDS = {
    'Battle name': 'A short name for this scenario.',
    'Your name': 'So we know whose expectations these are.',
    'Date and place': 'Optional, e.g. 2 July 1863, Gettysburg.',
    'Side A': 'Name of the first side, e.g. Union.',
    'Side B': 'Name of the second side, e.g. Confederate.',
    'Source': 'Optional: the historical battle or book this is based on.',
    'Who should hold the field?': 'A, B (or the side name) or Nobody. Leave blank if you have no view.',
    'Expected casualties, side A': 'Killed + wounded + captured, share of side A\'s men: e.g. 10-25%, <10%, >40%.',
    'Expected casualties, side B': 'The same for side B.',
    'Expected casualties, winner': 'Or say it for whichever side wins.',
    'Expected casualties, loser': 'And for whichever side loses.',
    'How sure are you?': 'low, medium or high.',
    'Notes': 'Anything else: what should happen and why.',
}
UNIT_COLUMNS = ('Unit id', 'Name', 'Side', 'Type', 'Subtype', 'Men', 'Guns', 'Experience', 'Morale', 'Weapon',
                'Melee only?', 'Should it break?')
ORDER_COLUMNS = ('Round', 'Minutes', 'Unit id', 'Action', 'Target', 'Ammo')
FEEDBACK_COLUMNS = ('Your name', 'Engine version', 'Casualties', 'Breaks', 'Right side favoured?', 'Comments')
FEEDBACK_KEYS = ('author', 'engine', 'casualties', 'breaks', 'right_side_favoured', 'comments')
REQUIRED_TABS = ('Battle', 'Units', 'Orders')

START_HERE = [
    'Battle scenario workbook',
    '',
    'Use this workbook to describe a battle, say what you expect to happen, and later tell us how the result felt.',
    'The engine runs your battle hundreds of times with different dice and reports how often it agrees with you.',
    '',
    'Battle tab: name the battle and the two sides, and (optionally) who should win and roughly what each side loses.',
    'Units tab: one row per regiment or battery. Men is the number of soldiers (for artillery: crew; leave Men and',
    '  Guns blank for a standard battery of 8 guns). Experience and Morale are 1-10, as on the unit\'s stat block.',
    '  "Should it break?": yes, no, or blank if you have no view.',
    'Orders tab: what each unit does, round by round. A round is a stretch of time (e.g. 60 or 90 minutes).',
    '  fire = shoot at the Target; charge = close and fight hand to hand; hold = stay put. Units without a row hold.',
    '  Artillery can pick Ammo: round shot (long range), shell, or canister (close range).',
    '  Not sure how to write the orders? Describe the plan in Notes on the Battle tab and we will write them.',
    'Feedback tab: after you have seen the engine\'s report, rate how it felt. Copy the engine version from the report.',
    '',
    'How to send it back: Google Sheets: File > Download > Microsoft Excel (.xlsx). Excel: just save. Send us the file.',
]


def _vocab() -> dict:
    return get_config()['sheets']


def _choices(words) -> str:
    words = list(words)
    return ', '.join(words[:-1]) + ' or ' + words[-1]


def _text(value) -> str | None:
    """A cell's text, trimmed; None when blank."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _word(value) -> str | None:
    text = _text(value)
    return None if text is None else text.lower()


def _whole(value) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def _number(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def _rows(ws, width: int):
    """(row number, values) for every non-blank row below the header."""
    for r in range(2, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, width + 1)]
        if any(_text(v) is not None for v in values):
            yield r, values


# =============================================================================
# Reading a filled-in workbook
# =============================================================================

def _read_battle(ws, errors: list) -> dict:
    fields = {}
    for r in range(2, ws.max_row + 1):
        label = _text(ws.cell(r, 1).value)
        if label in BATTLE_FIELDS:
            fields[label] = ws.cell(r, 2).value
    return fields


def _side_of(value, sides: list) -> int | None:
    word = _word(value)
    if word in ('a', 'b'):
        return 'ab'.index(word)
    for i, name in enumerate(sides):
        if word == name.lower():
            return i
    return None


def _read_units(ws, sides: list, errors: list) -> tuple:
    vocab = _vocab()
    types, weapons = vocab['unit_types'], vocab['weapons']
    units, should_break, first_row = [], {}, {}
    for r, (uid, name, side, kind, subtype, men, guns, xp, morale, weapon, melee, breaks) in _rows(ws, 12):
        where = f'Units row {r}'
        problems = []
        uid = _text(uid)
        if uid is None:
            problems.append('Unit id is missing.')
        elif uid in first_row:
            problems.append(f"Unit id '{uid}' is already used on row {first_row[uid]}.")
        else:
            first_row[uid] = r
        if _word(kind) not in types:
            problems.append(f'Type must be {_choices(types)}; got {kind!r}.')
        side_no = _side_of(side, sides)
        if side_no is None:
            problems.append(f'Side must be A or B; got {side!r}.')
        is_art = types.get(_word(kind)) == 'art'
        size = _whole(men)
        if men is not None and (size is None or size <= 0):
            problems.append(f'Men must be a whole number above 0; got {men!r}.')
        elif men is None and not is_art:
            problems.append('Men is missing.')
        gun_count = _whole(guns)
        if guns is not None and (gun_count is None or gun_count < 0):
            problems.append(f'Guns must be a whole number, 0 or more; got {guns!r}.')
        for label, value in (('Experience', xp), ('Morale', morale)):
            v = _whole(value)
            if v is None or not 1 <= v <= 10:
                problems.append(f'{label} must be a whole number from 1 to 10; got {value!r}.')
        if _word(weapon) not in weapons:
            problems.append(f"Weapon must be one of {', '.join(weapons)}; got {weapon!r}.")
        if _word(melee) not in (None, 'yes', 'no'):
            problems.append(f'Melee only? must be yes or no; got {melee!r}.')
        if _word(breaks) not in (None, 'yes', 'no'):
            problems.append(f'Should it break? must be yes, no or blank; got {breaks!r}.')
        if problems:
            errors.extend(f'{where}: {p}' for p in problems)
            continue
        unit = {'id': uid, 'name': _text(name) or uid, 'side': sides[side_no], 'type': types[_word(kind)],
                'subtype': _word(subtype)}
        if size is not None:
            unit['size'] = size
        if gun_count is not None:
            unit['guns'] = gun_count
        unit['stats'] = f"{_whole(xp)}/{_whole(morale)}/{weapons[_word(weapon)]}/{1 if _word(melee) == 'yes' else 0}"
        units.append(unit)
        if _word(breaks) is not None:
            should_break[uid] = _word(breaks) == 'yes'
    return units, should_break, set(first_row)


def _read_orders(ws, unit_ids: set, errors: list) -> list:
    vocab = _vocab()
    actions, ammo_words = vocab['actions'], vocab['ammo']
    rounds: dict = {}
    durations: dict = {}        # round -> minutes, from the first row that gives both
    for r, (rnd, minutes, uid, action, target, ammo) in _rows(ws, 6):
        where = f'Orders row {r}'
        problems = []
        number = _whole(rnd)
        if number is None or number <= 0:
            problems.append(f'Round must be a whole number above 0; got {rnd!r}.')
        length = _number(minutes) if minutes is not None else None
        if minutes is None:
            problems.append('Minutes must be a number above 0.')
        elif length is None or length <= 0:
            problems.append(f'Minutes must be a number above 0; got {minutes!r}.')
        elif number is not None and number > 0 and durations.setdefault(number, length) != length:
            problems.append(f'round {number} is {durations[number]:g} minutes on an earlier row; got {minutes:g}.')
        uid = _text(uid)
        if uid not in unit_ids:
            problems.append(f"Unit {uid!r} isn't on the Units tab.")
        word = _word(action)
        if word not in actions:
            problems.append(f'Action must be {_choices(actions)}; got {action!r}.')
        target_id = _text(target)
        if target_id is not None and target_id not in unit_ids and word != 'hold':
            problems.append(f"Target '{target_id}' isn't on the Units tab.")
        elif target_id is None and word in ('fire', 'charge'):
            problems.append(f'a {word} needs a Target.')
        if _word(ammo) is not None and _word(ammo) not in ammo_words:
            problems.append(f'Ammo must be {_choices(ammo_words)}; got {ammo!r}.')
        if number is not None and length is not None and length > 0 and number in rounds \
                and uid in rounds[number]['orders']:
            problems.append(f'{uid} already has an order in round {number}.')
        if problems:
            errors.extend(f'{where}: {p}' for p in problems)
            continue
        entry = rounds.setdefault(number, {'duration': length, 'orders': {}})
        if word != 'hold':
            order = {'target': target_id, 'mode': actions[word]}
            if _word(ammo) is not None:
                order['ammo'] = ammo_words[_word(ammo)]
            entry['orders'][uid] = order
    return [{'duration': int(v['duration']) if v['duration'].is_integer() else v['duration'], 'orders': v['orders']}
            for _, v in sorted(rounds.items())]


def _read_feedback(ws, errors: list) -> list:
    scales = _vocab()['feedback']
    entries = []
    for r, values in _rows(ws, 6):
        entry = dict(zip(FEEDBACK_KEYS, values))
        problems = []
        for key, label in (('casualties', 'Casualties'), ('breaks', 'Breaks'),
                           ('right_side_favoured', 'Right side favoured?')):
            word = _word(entry[key])
            if word is not None and word not in scales[key]:
                problems.append(f'{label} must be {_choices(scales[key])}; got {entry[key]!r}.')
            entry[key] = word
        if problems:
            errors.extend(f'Feedback row {r}: {p}' for p in problems)
            continue
        entry['author'], entry['engine'], entry['comments'] = (_text(entry['author']), _text(entry['engine']),
                                                               _text(entry['comments']))
        entries.append(entry)
    return entries


def read_workbook(path) -> dict:
    """
    A scenario from a filled-in workbook.

    Raises
    ------
    SetupError
        Listing every problem, each with its tab and row; then anything the engine refuses.
    """
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError):
        raise SetupError([f"{path.name} isn't an Excel workbook (.xlsx)."]) from None
    missing = [t for t in REQUIRED_TABS if t not in wb.sheetnames]
    if missing:
        raise SetupError([f"{path.name} is missing the tab(s) {', '.join(missing)}: start from the template."])

    errors: list[str] = []
    battle = _read_battle(wb['Battle'], errors)
    sides = [_text(battle.get('Side A')) or 'Side A', _text(battle.get('Side B')) or 'Side B']

    confidence = _word(battle.get('How sure are you?'))
    if confidence is not None and confidence not in validation.CONFIDENCE:
        errors.append(f"Battle: 'How sure are you?' must be low, medium or high; got {battle['How sure are you?']!r}.")
    casualties = {}
    for label, key in (('side A', sides[0]), ('side B', sides[1]), ('winner', 'winner'), ('loser', 'loser')):
        field = f'Expected casualties, {label}'
        band = _text(battle.get(field))
        if band is None:
            continue
        try:
            validation.parse_band(band)
            casualties[key] = band
        except ValueError:
            errors.append(f"Battle: '{field}' must be a band like 25-35%, <10% or >40%; got {battle[field]!r}.")
    field_held = None
    holder = _text(battle.get('Who should hold the field?'))
    if holder is not None:
        if holder.lower() in ('nobody', 'none'):
            field_held = 'none'
        elif (side := _side_of(holder, sides)) is not None:
            field_held = sides[side]
        else:
            errors.append(f"Battle: 'Who should hold the field?' must be A, B, a side name or nobody; "
                          f'got {holder!r}.')

    units, should_break, unit_ids = _read_units(wb['Units'], sides, errors)
    rounds = _read_orders(wb['Orders'], unit_ids, errors)
    feedback = _read_feedback(wb['Feedback'], errors) if 'Feedback' in wb.sheetnames else []
    if errors:
        raise SetupError(errors)

    meta = {'name': _text(battle.get('Battle name')) or path.stem}
    if _text(battle.get('Date and place')):
        meta['date'] = _text(battle['Date and place'])
    scenario = {'meta': meta, 'sides': sides, 'units': units, 'rounds': rounds, 'finish': True}
    notes = _text(battle.get('Notes'))
    breaks = [u for u, b in should_break.items() if b]
    holds = [u for u, b in should_break.items() if not b]
    if field_held is not None or breaks or holds or casualties:
        expect = {'author': _text(battle.get('Your name')), 'source': _text(battle.get('Source')),
                  'confidence': confidence, 'notes': notes}
        if field_held is not None:
            expect['field_held_by'] = field_held
        expect.update({'breaks': breaks, 'holds': holds, 'casualties': casualties})
        scenario['expect'] = expect
    elif notes:
        meta['notes'] = notes
    if feedback:
        scenario['feedback'] = feedback
    meta['imported_from'] = path.name

    # what the engine and the checker would refuse (rows and rounds are well formed by now)
    battle_from_setup(scenario, seed=0)
    validation.read_expect(scenario)
    return scenario


# =============================================================================
# Writing a template (blank, or filled from a scenario)
# =============================================================================

_HEADER = Font(bold=True)
_HEADER_FILL = PatternFill('solid', fgColor='E8E6E1')


def _style_header(ws, widths: list) -> None:
    for i, width in enumerate(widths, start=1):
        cell = ws.cell(1, i)
        cell.font, cell.fill = _HEADER, _HEADER_FILL
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.freeze_panes = 'A2'


def _dropdown(ws, column: str, words, rows: int) -> None:
    dv = DataValidation(type='list', formula1='"' + ','.join(words) + '"', allow_blank=True)
    dv.add(f'{column}2:{column}{rows}')
    ws.add_data_validation(dv)


def _side_letter(side, sides) -> str:
    if isinstance(side, int):
        return 'AB'[side]
    return 'AB'[sides.index(side)]


def _fill_from(wb, scenario: dict) -> None:
    vocab = _vocab()
    types = {code: word for word, code in vocab['unit_types'].items()}
    weapons = {code: word for word, code in vocab['weapons'].items()}
    actions = {mode: word for word, mode in vocab['actions'].items()}
    ammo_words = {code: word for word, code in vocab['ammo'].items()}
    sides = scenario.get('sides') or ['Side A', 'Side B']
    expect = scenario.get('expect') or {}
    meta = scenario.get('meta') or {}

    def side_name(value):
        return sides[value] if isinstance(value, int) else value

    values = {
        'Battle name': meta.get('name'), 'Your name': expect.get('author'), 'Date and place': meta.get('date'),
        'Side A': sides[0], 'Side B': sides[1], 'Source': expect.get('source'),
        'How sure are you?': expect.get('confidence'), 'Notes': expect.get('notes') or meta.get('notes'),
    }
    if 'field_held_by' in expect:
        held = expect['field_held_by']
        values['Who should hold the field?'] = 'Nobody' if held in (None, 'none') else side_name(held)
    for key, band in (expect.get('casualties') or {}).items():
        label = key if key in ('winner', 'loser') else ('side A' if side_name(key) == sides[0] else 'side B')
        values[f'Expected casualties, {label}'] = band
    ws = wb['Battle']
    for r in range(2, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if values.get(label) is not None:
            ws.cell(r, 2).value = values[label]

    breaks, holds = set(expect.get('breaks') or []), set(expect.get('holds') or [])
    ws = wb['Units']
    for r, u in enumerate(scenario['units'], start=2):
        uid = u.get('id', u.get('name'))
        xp, morale, weapon, melee = (int(x) for x in u['stats'].split('/'))
        kind = u.get('type')
        row = [uid, u.get('name', uid), _side_letter(u['side'], sides), types[kind or 'inf'],
               u.get('subtype') or 'line', u.get('size'), u.get('guns'), xp, morale, weapons[weapon],
               'yes' if melee else 'no', 'yes' if uid in breaks else 'no' if uid in holds else None]
        for c, value in enumerate(row, start=1):
            ws.cell(r, c).value = value

    ws = wb['Orders']
    r = 2
    first = scenario['units'][0].get('id', scenario['units'][0].get('name'))
    for n, rnd in enumerate(scenario.get('rounds', []), start=1):
        orders = rnd.get('orders') or {}
        rows = [[n, rnd['duration'], uid, actions[o.get('mode', 'idle')], o.get('target'),
                 ammo_words.get(o.get('ammo'))] for uid, o in orders.items()]
        for row in rows or [[n, rnd['duration'], first, 'hold', None, None]]:
            for c, value in enumerate(row, start=1):
                wb['Orders'].cell(r, c).value = value
            r += 1

    ws = wb['Feedback']
    for r, entry in enumerate(scenario.get('feedback') or [], start=2):
        for c, key in enumerate(FEEDBACK_KEYS, start=1):
            ws.cell(r, c).value = entry.get(key)


def write_template(path, scenario: dict | None = None) -> None:
    """Save a workbook for players: blank, or filled from ``scenario`` for them to adjust."""
    vocab = _vocab()
    rows = vocab['rows']
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Start here'
    for r, line in enumerate(START_HERE, start=1):
        ws.cell(r, 1).value = line
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 120

    ws = wb.create_sheet('Battle')
    for c, title in enumerate(('Field', 'Value', 'Help'), start=1):
        ws.cell(1, c).value = title
    for r, (field, help_text) in enumerate(BATTLE_FIELDS.items(), start=2):
        ws.cell(r, 1).value = field
        ws.cell(r, 3).value = help_text
        ws.cell(r, 3).alignment = Alignment(wrap_text=True)
    _style_header(ws, [30, 30, 70])
    dv = DataValidation(type='list', formula1='"' + ','.join(validation.CONFIDENCE) + '"', allow_blank=True)
    dv.add(f'B{2 + list(BATTLE_FIELDS).index("How sure are you?")}')
    ws.add_data_validation(dv)

    ws = wb.create_sheet('Units')
    for c, title in enumerate(UNIT_COLUMNS, start=1):
        ws.cell(1, c).value = title
    _style_header(ws, [12, 24, 7, 11, 15, 8, 7, 11, 9, 18, 12, 15])
    subtypes = list(dict.fromkeys(s for kind in UNIT_SUBTYPES.values() for s in sorted(kind)))
    _dropdown(ws, 'C', ['A', 'B'], rows)
    _dropdown(ws, 'D', vocab['unit_types'], rows)
    _dropdown(ws, 'E', subtypes, rows)
    _dropdown(ws, 'J', vocab['weapons'], rows)
    _dropdown(ws, 'K', ['yes', 'no'], rows)
    _dropdown(ws, 'L', ['yes', 'no'], rows)

    ws = wb.create_sheet('Orders')
    for c, title in enumerate(ORDER_COLUMNS, start=1):
        ws.cell(1, c).value = title
    _style_header(ws, [8, 9, 12, 10, 12, 12])
    _dropdown(ws, 'D', vocab['actions'], rows)
    _dropdown(ws, 'F', vocab['ammo'], rows)

    ws = wb.create_sheet('Feedback')
    for c, title in enumerate(FEEDBACK_COLUMNS, start=1):
        ws.cell(1, c).value = title
    _style_header(ws, [16, 15, 15, 15, 20, 60])
    scales = vocab['feedback']
    _dropdown(ws, 'C', scales['casualties'], rows)
    _dropdown(ws, 'D', scales['breaks'], rows)
    _dropdown(ws, 'E', scales['right_side_favoured'], rows)

    if scenario is not None:
        _fill_from(wb, scenario)
    wb.save(path)

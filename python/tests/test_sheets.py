import os

import openpyxl
import pytest

from imperial_generals import api, sheets, validation
from imperial_generals.battles.setup import SetupError, load_setup_file

EXAMPLE = os.path.join(os.path.dirname(__file__), '../../scenarios/example_ridge.yaml')
TABS = ['Start here', 'Battle', 'Units', 'Orders', 'Feedback']


def fill(path, battle=None, units=(), orders=(), feedback=()):
    """Fill a fresh template the way a player would."""
    sheets.write_template(path)
    wb = openpyxl.load_workbook(path)
    for field, value in (battle or {}).items():
        ws = wb['Battle']
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == field)
        ws.cell(row, 2).value = value
    for tab, rows in (('Units', units), ('Orders', orders), ('Feedback', feedback)):
        for i, values in enumerate(rows, start=2):
            for j, value in enumerate(values, start=1):
                wb[tab].cell(i, j).value = value
    wb.save(path)
    return path


BATTLE = {
    'Battle name': 'Cemetery Ridge', 'Your name': 'Jane', 'Date and place': '2 July 1863, Gettysburg',
    'Side A': 'Union', 'Side B': 'Confederate', 'Source': 'Fox 1889',
    'Who should hold the field?': 'Union', 'Expected casualties, side A': '10-25%',
    'Expected casualties, side B': '>30%', 'How sure are you?': 'Medium', 'Notes': 'the charge fails',
}
UNITS = [
    # id, name, side, type, subtype, men, guns, experience, morale, weapon, melee only, should it break
    ('us_1mn', '1st Minnesota', 'A', 'infantry', 'line', 900, None, 6, 6, 'rifled musket', 'no', 'no'),
    ('us_bat', 'Battery B', 'A', 'artillery', 'battery', None, None, 5, 5, 'smoothbore musket', 'no', None),
    ('cs_8al', '8th Alabama', 'B', 'infantry', 'line', 1000, None, 5, 6, 'smoothbore musket', 'no', 'yes'),
    ('cs_cav', 'Stuart', 'B', 'cavalry', 'heavy', 400, None, 6, 6, 'smoothbore musket', 'yes', None),
    ('cs_bat', 'Ross', 'B', 'artillery', 'horse battery', 50, 4, 4, 4, 'smoothbore musket', 'no', None),
]
ORDERS = [
    # round, minutes, unit, action, target, ammo
    (1, 60, 'us_1mn', 'fire', 'cs_8al', None),
    (1, 60, 'cs_8al', 'fire', 'us_1mn', None),
    (1, 60, 'us_bat', 'fire', 'cs_8al', 'shell'),
    (1, 60, 'cs_cav', 'hold', None, None),
    (2, 30, 'cs_8al', 'charge', 'us_1mn', None),
    (2, 30, 'us_bat', 'fire', 'cs_8al', 'canister'),
]
FEEDBACK = [
    # name, engine version, casualties, breaks, right side favoured?, comments
    ('Jane', 'abc12345', 'too heavy', 'about right', 'yes', 'bloody'),
    ('Sam', 'abc12345', 'about right', 'too early', 'unsure', None),
]


def test_template_has_tabs_headers_and_dropdowns(tmp_path):
    sheets.write_template(tmp_path / 't.xlsx')
    wb = openpyxl.load_workbook(tmp_path / 't.xlsx')
    assert wb.sheetnames == TABS
    assert [c.value for c in wb['Units'][1]] == list(sheets.UNIT_COLUMNS)
    assert [c.value for c in wb['Orders'][1]] == list(sheets.ORDER_COLUMNS)
    assert [c.value for c in wb['Feedback'][1]] == list(sheets.FEEDBACK_COLUMNS)
    assert [wb['Battle'].cell(r, 1).value for r in range(2, 2 + len(sheets.BATTLE_FIELDS))] == list(sheets.BATTLE_FIELDS)
    lists = {str(dv.sqref)[:1]: dv.formula1 for dv in wb['Units'].data_validations.dataValidation}
    assert 'infantry' in lists['D'] and 'line' in lists['E'] and 'rifled musket' in lists['J']
    orders = [dv.formula1 for dv in wb['Orders'].data_validations.dataValidation]
    assert any('charge' in f for f in orders) and any('canister' in f for f in orders)
    assert 'how to send it back' in ' '.join(str(c.value) for row in wb['Start here'].iter_rows() for c in row).lower()


def test_import_turns_a_filled_sheet_into_a_scenario(tmp_path):
    path = fill(tmp_path / 'b.xlsx', BATTLE, UNITS, ORDERS, FEEDBACK)
    sc = sheets.read_workbook(path)
    assert sc['meta'] == {'name': 'Cemetery Ridge', 'date': '2 July 1863, Gettysburg', 'imported_from': 'b.xlsx'}
    assert sc['sides'] == ['Union', 'Confederate']
    us, bat, cs, cav, ross = sc['units']
    assert us == {'id': 'us_1mn', 'name': '1st Minnesota', 'side': 'Union', 'type': 'inf', 'subtype': 'line',
                  'size': 900, 'stats': '6/6/1/0'}
    assert bat == {'id': 'us_bat', 'name': 'Battery B', 'side': 'Union', 'type': 'art', 'subtype': 'battery',
                   'stats': '5/5/0/0'}                                       # standard battery
    assert cav['stats'] == '6/6/0/1' and ross['guns'] == 4 and ross['size'] == 50
    assert sc['rounds'] == [
        {'duration': 60, 'orders': {'us_1mn': {'target': 'cs_8al', 'mode': 'ranged'},
                                    'cs_8al': {'target': 'us_1mn', 'mode': 'ranged'},
                                    'us_bat': {'target': 'cs_8al', 'mode': 'ranged', 'ammo': 'shell'}}},
        {'duration': 30, 'orders': {'cs_8al': {'target': 'us_1mn', 'mode': 'melee'},
                                    'us_bat': {'target': 'cs_8al', 'mode': 'ranged', 'ammo': 'canister'}}},
    ]
    assert sc['finish'] is True
    assert sc['expect'] == {
        'author': 'Jane', 'source': 'Fox 1889', 'confidence': 'medium', 'notes': 'the charge fails',
        'field_held_by': 'Union', 'breaks': ['cs_8al'], 'holds': ['us_1mn'],
        'casualties': {'Union': '10-25%', 'Confederate': '>30%'},
    }
    assert sc['feedback'] == [
        {'author': 'Jane', 'engine': 'abc12345', 'casualties': 'too heavy', 'breaks': 'about right',
         'right_side_favoured': 'yes', 'comments': 'bloody'},
        {'author': 'Sam', 'engine': 'abc12345', 'casualties': 'about right', 'breaks': 'too early',
         'right_side_favoured': 'unsure', 'comments': None},
    ]
    # it runs and can be checked
    assert validation.check(sc, seeds=2)['total'] == 5


def test_blank_expectations_are_left_out(tmp_path):
    battle = {'Battle name': 'Quiet', 'Side A': 'Blue', 'Side B': 'Red'}
    units = [('b', 'B', 'A', 'infantry', 'line', 100, None, 5, 5, 'smoothbore musket', 'no', None),
             ('r', 'R', 'B', 'infantry', 'line', 100, None, 5, 5, 'smoothbore musket', None, None)]
    sc = sheets.read_workbook(fill(tmp_path / 'q.xlsx', battle, units))
    assert 'expect' not in sc and 'feedback' not in sc and sc['rounds'] == []
    assert sc['meta'] == {'name': 'Quiet', 'imported_from': 'q.xlsx'}
    assert sc['units'][1]['stats'] == '5/5/0/0'


def test_unnamed_sides_and_field_by_letter(tmp_path):
    units = [('b', None, 'A', 'infantry', 'line', 100, None, 5, 5, 'smoothbore musket', 'no', None),
             ('r', None, 'B', 'infantry', 'line', 100, None, 5, 5, 'smoothbore musket', 'no', None)]
    sc = sheets.read_workbook(fill(tmp_path / 'n.xlsx', {'Who should hold the field?': 'B'}, units))
    assert sc['sides'] == ['Side A', 'Side B'] and sc['meta']['name'] == 'n'
    assert sc['units'][0]['name'] == 'b' and sc['units'][0]['side'] == 'Side A'
    assert sc['expect']['field_held_by'] == 'Side B'
    sc = sheets.read_workbook(fill(tmp_path / 'm.xlsx', {'Who should hold the field?': 'Nobody'}, units))
    assert sc['expect']['field_held_by'] == 'none'


def test_problems_point_to_the_sheet_and_row(tmp_path):
    units = [
        ('a', 'A', 'A', 'horse', 'line', 100, None, 5, 5, 'smoothbore musket', 'no', None),
        ('b', 'B', 'C', 'infantry', 'line', 'many', None, 11, 5, 'laser', 'maybe', 'perhaps'),
        ('a', 'A2', 'B', 'infantry', 'line', 100, None, 5, 5, 'smoothbore musket', 'no', None),
        (None, 'nameless', 'B', 'infantry', 'line', 100, None, 5, 5, 'smoothbore musket', 'no', None),
    ]
    orders = [
        (1, 60, 'a', 'shoot', 'zz', None),
        ('two', 60, 'b', 'fire', 'a', None),
        (1, 30, 'b', 'fire', 'a', 'grape'),
        (2, None, 'a', 'charge', None, None),
    ]
    battle = {'How sure are you?': 'certain', 'Expected casualties, side A': 'lots'}
    with pytest.raises(SetupError) as info:
        sheets.read_workbook(fill(tmp_path / 'bad.xlsx', battle, units, orders,
                                  [('Jane', 'x', 'way too heavy', None, None, None)]))
    errors = info.value.errors
    assert errors == [
        "Battle: 'How sure are you?' must be low, medium or high; got 'certain'.",
        "Battle: 'Expected casualties, side A' must be a band like 25-35%, <10% or >40%; got 'lots'.",
        "Units row 2: Type must be infantry, cavalry or artillery; got 'horse'.",
        "Units row 3: Side must be A or B; got 'C'.",
        "Units row 3: Men must be a whole number above 0; got 'many'.",
        "Units row 3: Experience must be a whole number from 1 to 10; got 11.",
        "Units row 3: Weapon must be one of unarmed / pikes, matchlock, smoothbore musket, rifled musket, "
        "needle rifle; got 'laser'.",
        "Units row 3: Melee only? must be yes or no; got 'maybe'.",
        "Units row 3: Should it break? must be yes, no or blank; got 'perhaps'.",
        "Units row 4: Unit id 'a' is already used on row 2.",
        "Units row 5: Unit id is missing.",
        "Orders row 2: Action must be fire, charge or hold; got 'shoot'.",
        "Orders row 2: Target 'zz' isn't on the Units tab.",
        "Orders row 3: Round must be a whole number above 0; got 'two'.",
        "Orders row 4: round 1 is 60 minutes on an earlier row; got 30.",
        "Orders row 4: Ammo must be round shot, shell or canister; got 'grape'.",
        "Orders row 5: Minutes must be a number above 0.",
        "Orders row 5: a charge needs a Target.",
        "Feedback row 2: Casualties must be far too light, too light, about right, too heavy or far too heavy; "
        "got 'way too heavy'.",
    ]


def test_import_checks_the_scenario_too(tmp_path):
    # rows are well formed, but the engine refuses the scenario: an infantry subtype on cavalry
    units = [('a', 'A', 'A', 'cavalry', 'line', 100, None, 5, 5, 'smoothbore musket', 'no', None),
             ('b', 'B', 'B', 'infantry', 'line', 100, None, 5, 5, 'smoothbore musket', 'no', None)]
    with pytest.raises(SetupError) as info:
        sheets.read_workbook(fill(tmp_path / 'x.xlsx', {}, units))
    assert any("subtype 'line' is not valid for cavalry" in e for e in info.value.errors)


def test_not_a_workbook_or_missing_tabs(tmp_path):
    (tmp_path / 'notes.xlsx').write_text('hello', encoding='utf-8')
    with pytest.raises(SetupError) as info:
        sheets.read_workbook(tmp_path / 'notes.xlsx')
    assert info.value.errors == ["notes.xlsx isn't an Excel workbook (.xlsx)."]
    wb = openpyxl.Workbook()
    wb.save(tmp_path / 'blank.xlsx')
    with pytest.raises(SetupError) as info:
        sheets.read_workbook(tmp_path / 'blank.xlsx')
    assert info.value.errors == ["blank.xlsx is missing the tab(s) Battle, Units, Orders: start from the template."]


# =============================================================================
# Export: send players a ready-made scenario to adjust
# =============================================================================

def test_export_then_import_gives_the_same_battle(tmp_path):
    original = load_setup_file(EXAMPLE)
    sheets.write_template(tmp_path / 'ridge.xlsx', scenario=original)
    back = sheets.read_workbook(tmp_path / 'ridge.xlsx')
    assert back['sides'] == original['sides'] and back['units'] == original['units']
    assert back['rounds'] == original['rounds']
    a = api.run_scenario({**original, 'finish': True}, seed=3)[2]
    b = api.run_scenario({**back, 'finish': True}, seed=3)[2]
    assert a == b


def test_export_carries_expectations_and_feedback(tmp_path):
    sc = sheets.read_workbook(fill(tmp_path / 'b.xlsx', BATTLE, UNITS, ORDERS, FEEDBACK))
    sheets.write_template(tmp_path / 'again.xlsx', scenario=sc)
    again = sheets.read_workbook(tmp_path / 'again.xlsx')
    assert again['expect'] == sc['expect'] and again['feedback'] == sc['feedback']
    assert again['units'] == sc['units'] and again['rounds'] == sc['rounds']


def test_export_needs_named_units(tmp_path):
    sc = {'units': [{'id': 'a', 'side': 0, 'size': 10, 'stats': '5/5/0/0'},
                    {'id': 'b', 'side': 1, 'size': 10, 'stats': '5/5/0/0'}],
          'rounds': [{'duration': 30}]}
    sheets.write_template(tmp_path / 'p.xlsx', scenario=sc)
    back = sheets.read_workbook(tmp_path / 'p.xlsx')
    assert back['units'][0]['type'] == 'inf' and back['units'][0]['subtype'] == 'line'   # generic -> line
    assert back['rounds'] == [{'duration': 30, 'orders': {}}]          # an idle round is kept as a 'hold' row


def test_numbers_typed_as_text_and_more_problems(tmp_path):
    units = [
        ('a', None, 'A', 'infantry', 'line', '900', None, 5.0, '5', 'smoothbore musket', 'no', None),   # fine
        ('b', None, 'B', 'infantry', 'line', None, None, True, 5, 'smoothbore musket', 'no', None),
        ('g', None, 'B', 'artillery', 'battery', 40, 'six', 5, 5, 'smoothbore musket', 'no', None),
    ]
    orders = [
        (1, '60', 'a', 'fire', 'b', None),                  # fine: minutes typed as text
        (1, 60, 'a', 'charge', 'b', None),                  # a second order for a in round 1
        (2, 'soon', 'a', 'fire', 'b', None),
        (2, True, 'a', 'fire', 'b', None),
        (3, 60, 'nobody', 'hold', None, None),
    ]
    with pytest.raises(SetupError) as info:
        sheets.read_workbook(fill(tmp_path / 'more.xlsx', {'Who should hold the field?': 'the French'}, units, orders))
    assert info.value.errors == [
        "Battle: 'Who should hold the field?' must be A, B, a side name or nobody; got 'the French'.",
        'Units row 3: Men is missing.',
        'Units row 3: Experience must be a whole number from 1 to 10; got True.',
        "Units row 4: Guns must be a whole number, 0 or more; got 'six'.",
        'Orders row 3: a already has an order in round 1.',
        "Orders row 4: Minutes must be a number above 0; got 'soon'.",
        'Orders row 5: Minutes must be a number above 0; got True.',
        "Orders row 6: Unit 'nobody' isn't on the Units tab.",
    ]


def test_whole_numbers_stored_as_floats():
    # some spreadsheet apps store whole numbers as 900.0
    assert sheets._whole(900.0) == 900 and sheets._whole(2.5) is None
